"""Optional-format document loaders with useful metadata and graceful degradation."""
from pathlib import Path
from .models import Document

SUPPORTED = {".pdf", ".docx", ".txt", ".md", ".csv"}

def load_document(path: str | Path) -> Document:
    p = Path(path); suffix = p.suffix.lower()
    if suffix not in SUPPORTED: raise ValueError(f"Unsupported document type: {suffix}")
    if suffix in {".txt", ".md"}:
        text = p.read_text(encoding="utf-8", errors="replace")
    elif suffix == ".csv":
        import csv
        with p.open(newline="", encoding="utf-8", errors="replace") as f:
            rows = list(csv.reader(f)); text = "\n".join(" | ".join(r) for r in rows)
    elif suffix == ".pdf":
        try:
            from pypdf import PdfReader
            pages = PdfReader(str(p)).pages
            text = "\n\n".join((page.extract_text() or "") for page in pages)
        except ImportError:
            try:
                import fitz
                doc = fitz.open(str(p)); text = "\n\n".join(x.get_text() for x in doc)
            except ImportError as e: raise RuntimeError("Install pypdf or pymupdf to read PDF files") from e
    else:
        try:
            from docx import Document as DocxDocument
            d = DocxDocument(str(p)); text = "\n".join(x.text for x in d.paragraphs)
            for table in d.tables: text += "\n" + "\n".join(" | ".join(c.text for c in row.cells) for row in table.rows)
        except ImportError as e: raise RuntimeError("Install python-docx to read DOCX files") from e
    stat = p.stat()
    return Document(text=text.strip(), source=str(p), metadata={"filename": p.name, "extension": suffix, "modified": stat.st_mtime, "bytes": stat.st_size})
